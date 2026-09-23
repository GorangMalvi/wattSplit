"""
One-time migration from the messy historical Excel workbook into the structured
backend/data/bills.xlsx workbook.

Usage:
    python -m app.migrate
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from dateutil import parser
from openpyxl import load_workbook

from . import db


SOURCE_PATH = Path(r"C:\Users\goran\Downloads\electricity golf 1.xlsx")

MONTH_MAP: Dict[str, str] = {
    "Apr": "2025-04",
    "May": "2025-05",
    "june": "2025-06",
    "July": "2025-07",
    "aug": "2025-08",
    "sept": "2025-09",
    "oct": "2025-10",
    "NOV": "2025-11",
    "dec": "2025-12",
    "JAN AFTER VIPUL": "2026-01",
    "FEB": "2026-02",
    "march": "2026-03",
    "april": "2026-04",
    "may26": "2026-05",
    "june26": "2026-06",
}

NAME_ALIASES = {
    "sak": "Gorang",
    "saksham": "Gorang",
    "ud": "Udbhav",
    "udbhav": "Udbhav",
    "udhbhav": "Udbhav",
    "vipul": "Akash",
    "nav": "Naveen",
    "naveen": "Naveen",
    "gorang": "Gorang",
    "akash": "Akash",
    "ujjwal": "Ujjwal",
    "ro": "ro",  # unrecognised abbreviation; kept verbatim
    "dg": "DG",
}


def normalize_name(raw: Any) -> str:
    if raw is None:
        return ""
    s = str(raw).strip().lower()
    return NAME_ALIASES.get(s, str(raw).strip())


def parse_date(v: Any) -> Optional[datetime]:
    """Parse a date value, repairing obvious typos like 31-09-2025 or year 2001."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        # Skip carry-forward / net-recharge labels.
        if "carry forward" in s.lower() or "net recharge" in s.lower():
            return None
        # Extract a date from strings like "(Actual)01-05-2025".
        m = re.search(r"(\d{1,2}[\-/]\d{1,2}[\-/]\d{4})", s)
        if m:
            s = m.group(1)
        # Repair month "012" -> "12" observed in one cell.
        parts = re.split(r"[\-/]", s)
        if len(parts) == 3:
            d, mo, y = parts[0], parts[1], parts[2]
            if mo.startswith("0") and len(mo) == 3:
                mo = mo[1:]
            s = f"{d}-{mo}-{y}"
        try:
            return parser.parse(s, dayfirst=True)
        except Exception:
            return None
    return None


def fix_typo_year(d: datetime, expected_year: int) -> datetime:
    """Correct obviously wrong years (e.g. 2001 when we expect 2026)."""
    if d.year < 2010 and expected_year >= 2020:
        return d.replace(year=expected_year)
    return d


def read_source_rows(sheet: str) -> List[Tuple[Any, ...]]:
    wb = load_workbook(SOURCE_PATH, data_only=True)
    return list(wb[sheet].iter_rows(values_only=True))


def find_left_header(rows: List[Tuple[Any, ...]]) -> Tuple[int, List[str]]:
    """Return (row_index, list_of_names) for the left-side meter table."""
    for i, row in enumerate(rows):
        if not row or row[0] is None:
            continue
        if str(row[0]).strip().lower() == "month":
            names = []
            for cell in row[1:]:
                if cell is None or str(cell).strip() == "":
                    break
                if str(cell).strip().lower() == "month":
                    break
                names.append(str(cell).strip())
            return i, names
    raise ValueError("No header row found")


def find_right_header(row: Tuple[Any, ...]) -> Tuple[Optional[int], List[str]]:
    """Return (start_column, list_of_names) for the right-side recharge table."""
    for idx, cell in enumerate(row):
        # Skip the left-hand table header; we want the second 'Month' cell.
        if idx == 0:
            continue
        if cell is not None and str(cell).strip().lower() == "month":
            names = []
            for c in row[idx + 1 :]:
                if c is None or str(c).strip() == "":
                    break
                if str(c).strip().lower() == "total":
                    break
                names.append(str(c).strip())
            return idx, names
    return None, []


def parse_month_sheet(sheet: str, code: str) -> Dict[str, Any]:
    rows = read_source_rows(sheet)
    header_idx, left_names = find_left_header(rows)
    expected_year = int(code[:4])

    # Locate important rows.
    units_idx = common_idx = total_idx = total_bill_idx = dg_idx = None
    for i, row in enumerate(rows):
        if not row or row[0] is None:
            continue
        label = str(row[0]).strip().lower()
        if "unit" in label and "consum" in label:
            units_idx = i
        elif label == "common unit":
            common_idx = i
        elif label == "total":
            total_idx = i
        elif "dg" in label:
            dg_idx = i
        elif "total bill" in label:
            total_bill_idx = i

    # Date rows in the left table.
    date_rows = []
    for i, row in enumerate(rows[header_idx + 1 :], start=header_idx + 1):
        d = parse_date(row[0])
        if d is not None:
            d = fix_typo_year(d, expected_year)
            date_rows.append((i, d, row))

    # ---- readings ----
    # Sub-units per name from the "Units consumed" row.
    sub_units: Dict[str, float] = {}
    if units_idx is not None:
        for j, name in enumerate(left_names):
            val = rows[units_idx][j + 1]
            if val is not None and isinstance(val, (int, float)):
                sub_units[normalize_name(name)] = float(val)

    # Determine start/end readings.
    start_readings: Dict[str, float] = {}
    end_readings: Dict[str, float] = {}

    if len(date_rows) >= 2:
        # Sort by date; use the last two rows.
        date_rows.sort(key=lambda x: x[1])
        _, _, start_row = date_rows[-2]
        _, _, end_row = date_rows[-1]
        for j, name in enumerate(left_names):
            n = normalize_name(name)
            v = start_row[j + 1]
            if v is not None and isinstance(v, (int, float)):
                start_readings[n] = float(v)
            v = end_row[j + 1]
            if v is not None and isinstance(v, (int, float)):
                end_readings[n] = float(v)
    elif len(date_rows) == 1:
        _, _, row = date_rows[0]
        for j, name in enumerate(left_names):
            n = normalize_name(name)
            v = row[j + 1]
            if v is not None and isinstance(v, (int, float)):
                # We cannot know if this is start or end without context.
                # Store as end; the migration loop will fill start from the
                # previous month when possible.
                end_readings[n] = float(v)

    # If sub-units are known but end readings are missing (e.g. Sept 2025),
    # we can compute the end readings later once the start readings are known.

    # ---- bills ----
    monthly_bill: Optional[float] = None
    if total_idx is not None:
        val = rows[total_idx][len(left_names) + 1]
        if val is not None and isinstance(val, (int, float)):
            monthly_bill = float(val)
    if monthly_bill is None and total_bill_idx is not None:
        val = rows[total_bill_idx][len(left_names) + 1]
        if val is not None and isinstance(val, (int, float)):
            monthly_bill = float(val)
    if monthly_bill is None and total_bill_idx is not None:
        # Fallback: sum the per-person values in the Total bill row.
        total = 0.0
        for j, _ in enumerate(left_names):
            val = rows[total_bill_idx][j + 1]
            if val is not None and isinstance(val, (int, float)):
                total += float(val)
        if total > 0:
            monthly_bill = total

    dg_bill = 0.0
    if dg_idx is not None:
        total = 0.0
        for j, _ in enumerate(left_names):
            val = rows[dg_idx][j + 1]
            if val is not None and isinstance(val, (int, float)):
                total += float(val)
        dg_bill = total

    # Compute main units as sum(sub-units) + sum(per-person common units).
    # In most sheets the last column of the "Common unit" row already equals
    # this total, but in Apr 2025 it only contains the common total.
    common_total = 0.0
    if common_idx is not None:
        for j, name in enumerate(left_names):
            n = normalize_name(name)
            val = rows[common_idx][j + 1]
            if val is not None and isinstance(val, (int, float)):
                common_total += float(val)
    sum_sub = sum(sub_units.values()) if sub_units else 0.0
    main_units = sum_sub + common_total

    recharges: List[Dict[str, Any]] = []

    # NOTE: The right-hand tables mix carry-forward balances, net-recharge totals
    # and actual payments in an ambiguous way. To avoid double-counting balances
    # as recharges we deliberately do NOT import the structured per-person
    # columns from those tables. Instead we rely on the explicit free-form
    # payment entries (date + amount + name) and the dedicated Sheet2 log.

    # ---- free-form recharges anywhere in the sheet ----
    for row in rows:
        cells = list(row)
        for i in range(len(cells) - 2):
            a, b, c = cells[i], cells[i + 1], cells[i + 2]
            # Look for a date + amount + name triple in any order.
            triples = [
                (a, b, c),
                (a, c, b),
                (b, a, c),
                (b, c, a),
                (c, a, b),
                (c, b, a),
            ]
            for d_val, amt_val, name_val in triples:
                d = parse_date(d_val)
                if d is None:
                    continue
                if not isinstance(amt_val, (int, float)):
                    continue
                if isinstance(amt_val, float) and pd.isna(amt_val):
                    continue
                if amt_val <= 0:
                    continue
                if name_val is None:
                    continue
                if isinstance(name_val, (int, float)):
                    continue
                name = normalize_name(name_val)
                if not name or name.lower() in ("month", "total", "reading"):
                    continue
                # Skip if this row is part of the left table header.
                if i == 0 and str(cells[0]).strip().lower() == "month":
                    continue
                recharges.append({
                    "date": d.date().isoformat(),
                    "roommate": name,
                    "amount": float(amt_val),
                    "notes": f"from sheet {sheet}",
                })
                break  # avoid duplicate triples from the same window

    return {
        "code": code,
        "sheet": sheet,
        "names": [normalize_name(n) for n in left_names],
        "start_readings": start_readings,
        "end_readings": end_readings,
        "sub_units": sub_units,
        "monthly_bill": monthly_bill,
        "dg_bill": dg_bill,
        "main_units": main_units,
        "recharges": recharges,
    }


def dedupe_recharges(recharges: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    seen = set()
    out = []
    for r in recharges:
        key = (r["date"], r["roommate"], round(r["amount"], 2))
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


def run_migration() -> None:
    print("Initializing backend database...")
    db.init_db(force=True)

    parsed_months = []
    for sheet, code in MONTH_MAP.items():
        print(f"Parsing sheet '{sheet}' -> {code} ...")
        parsed = parse_month_sheet(sheet, code)
        parsed_months.append(parsed)

    # Sort chronologically.
    parsed_months.sort(key=lambda x: x["code"])

    # ---- roommates ----
    all_names = set()
    name_first_month: Dict[str, str] = {}
    name_last_month: Dict[str, str] = {}
    for p in parsed_months:
        for name in p["names"]:
            if not name:
                continue
            all_names.add(name)
            if name not in name_first_month:
                name_first_month[name] = p["code"]
            name_last_month[name] = p["code"]

    # Also consider recharges.
    for p in parsed_months:
        for r in p["recharges"]:
            n = r["roommate"]
            if not n or n == "DG":
                continue
            all_names.add(n)
            if n not in name_first_month:
                name_first_month[n] = p["code"]
            name_last_month[n] = p["code"]

    latest_month = parsed_months[-1]["code"]
    roommate_ids = {}
    for name in sorted(all_names):
        is_active = name_last_month.get(name, "") == latest_month
        rid = db.add_roommate(
            name=name,
            join_date=name_first_month.get(name, ""),
            leave_date=name_last_month.get(name, "") if not is_active else "",
            is_active=is_active,
        )
        roommate_ids[name] = rid

    # ---- months & readings ----
    previous_end: Dict[str, float] = {}
    for idx, p in enumerate(parsed_months):
        code = p["code"]
        names = p["names"]

        # Resolve start/end readings.
        end_readings = {n: v for n, v in p["end_readings"].items()}
        start_readings = {n: v for n, v in p["start_readings"].items()}

        for n in names:
            if n not in start_readings:
                if n in previous_end:
                    start_readings[n] = previous_end[n]
                elif n in end_readings and n in p["sub_units"]:
                    # Compute start from end - sub_units.
                    start_readings[n] = end_readings[n] - p["sub_units"].get(n, 0)

        for n in names:
            if n not in end_readings:
                if n in start_readings and n in p["sub_units"]:
                    end_readings[n] = start_readings[n] + p["sub_units"].get(n, 0)
                elif n in previous_end:
                    end_readings[n] = previous_end[n]

        # Some sheets only record the *start* reading (e.g. Sept 2025).
        # If the parsed end equals the start but there is positive consumption,
        # compute the real end reading from the sub-units row.
        for n in names:
            end_val = end_readings.get(n)
            start_val = start_readings.get(n)
            if end_val is None or start_val is None:
                continue
            if abs(end_val - start_val) < 0.01:
                sub = p["sub_units"].get(n, 0)
                if sub > 0:
                    end_readings[n] = start_val + sub

        # Insert a base month before the very first billed month so that
        # the first month's sub-meter consumption can be computed.
        if idx == 0:
            base_code = db.get_previous_month(code) or ""
            if not base_code:
                year, mon = int(code[:4]), int(code[5:7])
                if mon == 1:
                    base_code = f"{year - 1}-12"
                else:
                    base_code = f"{year}-{mon - 1:02d}"
            db.add_month(
                month=base_code,
                monthly_bill=0.0,
                dg_bill=0.0,
                main_start_reading=0.0,
                main_end_reading=0.0,
                notes="Base readings before first billed month (auto-generated)",
            )
            for n in names:
                val = start_readings.get(n)
                if val is not None:
                    db.set_reading(base_code, n, val)

        # Main readings: infer sequentially from main_units.
        main_units = p["main_units"] or 0.0
        prev_main_end = 0.0
        prev_code = db.get_previous_month(code)
        if prev_code:
            prev_row = db.get_month(prev_code)
            if prev_row is not None:
                prev_main_end = float(prev_row["main_end_reading"]) if pd.notna(prev_row["main_end_reading"]) else 0.0
        main_start = prev_main_end
        main_end = main_start + main_units

        db.add_month(
            month=code,
            monthly_bill=p["monthly_bill"] or 0.0,
            dg_bill=p["dg_bill"] or 0.0,
            main_start_reading=main_start,
            main_end_reading=main_end,
            notes=f"Imported from sheet '{p['sheet']}'",
        )

        for n in names:
            val = end_readings.get(n)
            if val is not None:
                db.set_reading(code, n, val)

        previous_end = {n: end_readings.get(n, previous_end.get(n)) for n in names}

    # ---- recharges ----
    all_recharges = []
    for p in parsed_months:
        all_recharges.extend(p["recharges"])

    # Also parse Sheet2 for any recharges not already captured.
    sheet2_rows = read_source_rows("Sheet2")
    for row in sheet2_rows:
        cells = list(row)
        for i in range(len(cells) - 2):
            a, b, c = cells[i], cells[i + 1], cells[i + 2]
            triples = [
                (a, b, c), (a, c, b), (b, a, c),
                (b, c, a), (c, a, b), (c, b, a),
            ]
            for d_val, amt_val, name_val in triples:
                d = parse_date(d_val)
                if d is None:
                    continue
                if not isinstance(amt_val, (int, float)):
                    continue
                if isinstance(amt_val, float) and pd.isna(amt_val):
                    continue
                if amt_val <= 0:
                    continue
                if name_val is None:
                    continue
                if isinstance(name_val, (int, float)):
                    continue
                name = normalize_name(name_val)
                if not name or name.lower() in ("month", "total", "reading", "gpay", "card"):
                    continue
                all_recharges.append({
                    "date": d.date().isoformat(),
                    "roommate": name,
                    "amount": float(amt_val),
                    "notes": "from Sheet2",
                })
                break

    all_recharges = dedupe_recharges(all_recharges)
    for r in all_recharges:
        db.add_recharge(r["date"], r["roommate"], r["amount"], r["notes"])

    print(f"Migration complete. Database: {db.DB_PATH}")
    print(f"  Months: {len(parsed_months)}")
    print(f"  Roommates: {len(roommate_ids)}")
    print(f"  Recharges: {len(all_recharges)}")


if __name__ == "__main__":
    run_migration()
