"""
Generate a human-readable Excel report for a single month.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import calc, db


EXPORTS_DIR = Path(__file__).resolve().parent.parent / "exports"


def _format_header(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="366092")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _format_money(cell):
    cell.number_format = "#,##0.00"


def generate_month_report(month: str) -> Path:
    """Generate ``exports/{month}_report.xlsx`` and return its path."""
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    out_path = EXPORTS_DIR / f"{month}_report.xlsx"

    result = calc.calculate_month(month)
    month_row = db.get_month(month)
    notes = str(month_row["notes"]) if month_row is not None and pd.notna(month_row.get("notes")) else ""

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Title
    ws.merge_cells("A1:H1")
    ws["A1"] = f"Electricity Bill Split Report — {month}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="left")

    # Summary block
    summary = [
        ("Monthly Bill (excl. DG)", result["summary"]["total_bill"] - result["summary"]["dg_bill"]),
        ("DG Bill", result["summary"]["dg_bill"]),
        ("Total Bill", result["summary"]["total_bill"]),
        ("Main Start Reading", month_row["main_start_reading"] if month_row is not None else 0),
        ("Main End Reading", month_row["main_end_reading"] if month_row is not None else 0),
        ("Main Units", result["summary"]["main_units"]),
        ("Common Units", result["summary"]["main_units"] - sum(r["sub_units"] for r in result["roommates"])),
        ("Rate", result["summary"]["rate_per_unit"]),
        ("Active Roommates", len(result["roommates"])),
    ]
    row = 3
    for label, value in summary:
        ws.cell(row, 1, label).font = Font(bold=True)
        cell = ws.cell(row, 2, value)
        if isinstance(value, float):
            _format_money(cell)
        row += 1

    if notes:
        ws.cell(row, 1, "Notes").font = Font(bold=True)
        ws.cell(row, 2, notes)
        row += 1

    # Per-roommate split table
    row += 1
    headers = [
        "Roommate",
        "Previous Reading",
        "Current Reading",
        "Sub Units",
        "Common Share",
        "Total Units",
        "Energy Charge",
        "DG Share",
        "Total Bill",
        "Month Recharges",
        "Balance",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row, col, h)
        _format_header(cell)
    header_row = row
    row += 1

    for r in result["roommates"]:
        vals = [
            r["name"],
            r["previous_reading"],
            r["current_reading"],
            r["sub_units"],
            r["common_share"],
            r["total_units"],
            r["energy_charge"],
            r["dg_share"],
            r["total_bill"],
            r["recharges"],
            r["balance"],
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row, col, v)
            cell.border = border
            if isinstance(v, float):
                _format_money(cell)
        row += 1

    # Running balances
    row += 1
    ws.cell(row, 1, "Running Balances (cumulative up to this month)").font = Font(bold=True, size=12)
    row += 1
    bal_headers = ["Roommate", "Cumulative Bill", "Cumulative Recharges", "Balance"]
    for col, h in enumerate(bal_headers, 1):
        cell = ws.cell(row, col, h)
        _format_header(cell)
    row += 1
    for b in result["running_balances"]:
        vals = [b["name"], b["total_bill_cumulative"], b["total_recharges_cumulative"], b["balance"]]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row, col, v)
            cell.border = border
            if isinstance(v, float):
                _format_money(cell)
        row += 1

    # Recharge log for the month
    row += 1
    ws.cell(row, 1, f"Recharges recorded in {month}").font = Font(bold=True, size=12)
    row += 1
    rec_headers = ["Date", "Roommate", "Amount", "Notes"]
    for col, h in enumerate(rec_headers, 1):
        cell = ws.cell(row, col, h)
        _format_header(cell)
    row += 1
    rec_df = db.get_recharges()
    month_start = f"{month}-01"
    month_end = calc._month_end_date(month)
    recs = rec_df[(rec_df["date"] >= month_start) & (rec_df["date"] <= month_end)]
    if recs.empty:
        ws.cell(row, 1, "No recharges recorded for this month.")
    else:
        for _, rec in recs.iterrows():
            ws.cell(row, 1, str(rec["date"]))
            ws.cell(row, 2, str(rec["roommate"]))
            amt = ws.cell(row, 3, float(rec["amount"]))
            _format_money(amt)
            ws.cell(row, 4, str(rec.get("notes", "")))
            row += 1

    # Auto-width columns (rough)
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    wb.save(out_path)
    return out_path
