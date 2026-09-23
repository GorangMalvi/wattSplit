import pandas as pd
import openpyxl

file_path = r"C:\Users\goran\Downloads\electricity golf 1.xlsx"

wb = openpyxl.load_workbook(file_path, data_only=False)
print("Sheet names:", wb.sheetnames)

for sheet in wb.sheetnames:
    print(f"\n=== SHEET: {sheet} ===")
    df = pd.read_excel(file_path, sheet_name=sheet)
    print(f"Shape: {df.shape}")
    print(df.head(30).to_string())
    print("\nColumns:", df.columns.tolist())
